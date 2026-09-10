"""
src/exporters/excel_exporter.py - Multi-Tab Excel & CSV Exporter
GraphOne / FrontierAtlas Intelligence Graph Ingestion Pipeline
"""

import os
import logging
from typing import Dict, Any, List
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.config import EXCEL_OUTPUT_PATH, OUTPUT_DIR

logger = logging.getLogger("ExcelExporter")

class DataExporter:
    @staticmethod
    def flatten_startup(record: Dict[str, Any]) -> Dict[str, Any]:
        source = record.get("source", {})
        content = record.get("content", {})
        data = content.get("data", {})
        return {
            "schemaVersion": record.get("schemaVersion", "1.0"),
            "recordType": record.get("recordType", "STARTUP"),
            "source.name": source.get("name", ""),
            "source.url": source.get("url", ""),
            "content.entityName": content.get("entityName", ""),
            "content.data.employeeCount": data.get("employeeCount", None),
            "content.data.employeeCountMethod": data.get("employeeCountMethod", "UNKNOWN"),
            "collectedAt": record.get("collectedAt", "")
        }

    @staticmethod
    def flatten_product(record: Dict[str, Any]) -> Dict[str, Any]:
        source = record.get("source", {})
        content = record.get("content", {})
        return {
            "schemaVersion": record.get("schemaVersion", "1.0"),
            "recordType": record.get("recordType", "PRODUCT"),
            "source.name": source.get("name", ""),
            "source.url": source.get("url", ""),
            "content.startupName": content.get("startupName", ""),
            "content.pricingModel": content.get("pricingModel", "UNKNOWN"),
            "content.pricingMethod": content.get("pricingMethod", "UNKNOWN"),
            "collectedAt": record.get("collectedAt", "")
        }

    @staticmethod
    def flatten_paper(record: Dict[str, Any]) -> Dict[str, Any]:
        source = record.get("source", {})
        content = record.get("content", {})
        authors = content.get("authors", [])
        authors_str = ", ".join(authors) if isinstance(authors, list) else str(authors)
        return {
            "schemaVersion": record.get("schemaVersion", "1.0"),
            "recordType": record.get("recordType", "RESEARCH_PAPER"),
            "source.name": source.get("name", ""),
            "source.url": source.get("url", ""),
            "content.title": content.get("title", ""),
            "content.authors": authors_str,
            "content.paper_url": content.get("paper_url", ""),
            "content.github_url": content.get("github_url", ""),
            "content.github_stars": content.get("github_stars", None),
            "content.published_date": content.get("published_date", ""),
            "collectedAt": record.get("collectedAt", "")
        }

    @staticmethod
    def flatten_job(record: Dict[str, Any]) -> Dict[str, Any]:
        source = record.get("source", {})
        content = record.get("content", {})
        return {
            "schemaVersion": record.get("schemaVersion", "1.0"),
            "recordType": record.get("recordType", "JOB"),
            "source.name": source.get("name", ""),
            "source.url": source.get("url", ""),
            "content.company": content.get("company", ""),
            "content.title": content.get("title", ""),
            "content.date": content.get("date", ""),
            "content.is_remote": content.get("is_remote", True),
            "content.role_family": content.get("role_family", "Engineering"),
            "collectedAt": record.get("collectedAt", "")
        }

    @staticmethod
    def flatten_news(record: Dict[str, Any]) -> Dict[str, Any]:
        source = record.get("source", {})
        content = record.get("content", {})
        return {
            "schemaVersion": record.get("schemaVersion", "1.0"),
            "recordType": record.get("recordType", "NEWS"),
            "source.name": source.get("name", ""),
            "source.url": source.get("url", ""),
            "content.title": content.get("title", ""),
            "content.summary": content.get("summary", ""),
            "content.published_date": content.get("published_date", ""),
            "collectedAt": record.get("collectedAt", "")
        }

    @classmethod
    def export(cls, datasets: Dict[str, List[Dict[str, Any]]], mapping_log: List[Dict[str, Any]]):
        """Export datasets across 6 sheets in Excel workbook and separate CSV files."""
        logger.info("Exporting ingested datasets to Excel workbook and CSV files...")

        # Flatten records
        df_startups = pd.DataFrame([cls.flatten_startup(r) for r in datasets.get("startups", [])])
        df_products = pd.DataFrame([cls.flatten_product(r) for r in datasets.get("products", [])])
        df_papers = pd.DataFrame([cls.flatten_paper(r) for r in datasets.get("papers", [])])
        df_jobs = pd.DataFrame([cls.flatten_job(r) for r in datasets.get("jobs", [])])
        df_news = pd.DataFrame([cls.flatten_news(r) for r in datasets.get("news", [])])
        df_mapping = pd.DataFrame(mapping_log)

        # Export individual CSV files
        df_startups.to_csv(OUTPUT_DIR / "startups.csv", index=False)
        df_products.to_csv(OUTPUT_DIR / "products.csv", index=False)
        df_papers.to_csv(OUTPUT_DIR / "research_papers.csv", index=False)
        df_jobs.to_csv(OUTPUT_DIR / "jobs.csv", index=False)
        df_news.to_csv(OUTPUT_DIR / "news.csv", index=False)
        df_mapping.to_csv(OUTPUT_DIR / "entity_mapping_log.csv", index=False)
        logger.info("CSV outputs written to outputs/ directory.")

        # Export Excel Workbook with 6 sheets
        with pd.ExcelWriter(EXCEL_OUTPUT_PATH, engine="openpyxl") as writer:
            df_startups.to_excel(writer, sheet_name="Startups", index=False)
            df_products.to_excel(writer, sheet_name="Products", index=False)
            df_papers.to_excel(writer, sheet_name="Research Papers", index=False)
            df_jobs.to_excel(writer, sheet_name="Jobs", index=False)
            df_news.to_excel(writer, sheet_name="News", index=False)
            df_mapping.to_excel(writer, sheet_name="Entity Mapping Log", index=False)

            # Apply professional header styling
            workbook = writer.book
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            align = Alignment(horizontal="left", vertical="center")

            for sheet in workbook.worksheets:
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = align
                # Auto-fit column widths
                for col in sheet.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        logger.info(f"Successfully generated Excel workbook at: {EXCEL_OUTPUT_PATH}")
